"""Build the Phase 1 project report."""
from __future__ import annotations

from pathlib import Path
from textwrap import wrap

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output"
OUT.mkdir(exist_ok=True)
DOCX = OUT / "Phase1_G3-10_22I-1319.docx"
ARCH = OUT / "gatekeeper_architecture.png"

BLACK = "000000"
WHITE = "FFFFFF"


def font(size=28, bold=False):
    candidates = [
        Path("C:/Windows/Fonts/aptos.ttf"),
        Path("C:/Windows/Fonts/calibri.ttf"),
        Path("C:/Windows/Fonts/arial.ttf"),
    ]
    if bold:
        candidates = [
            Path("C:/Windows/Fonts/aptos-bold.ttf"),
            Path("C:/Windows/Fonts/calibrib.ttf"),
            Path("C:/Windows/Fonts/arialbd.ttf"),
        ] + candidates
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def make_architecture():
    canvas = Image.new("RGB", (2400, 1380), "white")
    d = ImageDraw.Draw(canvas)
    title = font(50, True)
    label = font(30, True)
    body = font(24)
    small = font(21)
    d.text((90, 45), "GateKeeper AI proposed CI/CD security architecture", fill="black", font=title)
    d.text((90, 108), "B0 and B1 are the baselines. G1, G2 and G3 mark the selected gaps.", fill="black", font=body)

    def box(x1, y1, x2, y2, heading, lines=(), width=3):
        d.rectangle((x1, y1, x2, y2), fill="white", outline="black", width=width)
        d.text((x1 + 24, y1 + 18), heading, fill="black", font=label)
        y = y1 + 64
        for line in lines:
            d.text((x1 + 24, y), line, fill="black", font=small)
            y += 31

    def arrow(x1, y1, x2, y2, width=4):
        d.line((x1, y1, x2, y2), fill="black", width=width)
        import math
        ang = math.atan2(y2 - y1, x2 - x1)
        length = 22
        for delta in (2.55, -2.55):
            d.line((x2, y2, x2 + length * math.cos(ang + delta), y2 + length * math.sin(ang + delta)), fill="black", width=width)

    box(90, 220, 400, 350, "Pull request", ["Changed Java code"])
    box(505, 220, 870, 350, "SAST scan", ["SpotBugs + FindSecBugs", "XML alert report"])
    box(980, 220, 1390, 350, "Normalise", ["Finding + ground truth", "stable structured contract"])
    arrow(400, 285, 505, 285)
    arrow(870, 285, 980, 285)

    box(90, 490, 570, 650, "B0 - raw SAST", ["Every alert reaches the gate", "No triage or explanation"])
    box(90, 740, 570, 900, "B1 - single LLM", ["One prompt with alert + file", "No repository tools or memory"])
    arrow(1180, 350, 1180, 420)
    arrow(1180, 420, 570, 570)
    arrow(1180, 420, 570, 820)

    d.rectangle((760, 440, 1880, 1070), fill="white", outline="black", width=5)
    d.text((800, 462), "Proposed evidence-grounded triage agent", fill="black", font=font(38, True))
    box(835, 560, 1240, 735, "LLM reasoner", ["Observe, choose tool", "or return verdict"])
    box(1365, 545, 1805, 755, "G1 - evidence", ["Read-only repository tools", "read_file / grep / definition", "Traceable code evidence"])
    box(835, 830, 1240, 1005, "G3 - budget", ["Steps, tokens and cost", "Early stop at fixed budget"])
    box(1365, 830, 1805, 1005, "G2 - abstention", ["Confidence threshold", "UNSURE to human review"])
    arrow(1240, 645, 1365, 645)
    arrow(1365, 690, 1240, 690)
    arrow(1038, 735, 1038, 830)
    arrow(1240, 700, 1365, 875)
    arrow(1390, 285, 1460, 440)

    box(1970, 530, 2310, 720, "Verdict", ["TP / FP / UNSURE", "confidence", "rationale + evidence"])
    box(1970, 835, 2310, 1015, "CI gate", ["TP blocks", "FP is suppressed", "UNSURE is reviewed"])
    arrow(1880, 700, 1970, 625)
    arrow(1880, 910, 1970, 925)
    arrow(2140, 720, 2140, 835)

    box(880, 1160, 1510, 1300, "Evaluation", ["Safety, precision, calibration, evidence, cost and latency"])
    arrow(2140, 1015, 1510, 1230)
    arrow(330, 650, 880, 1195, 3)
    arrow(330, 900, 880, 1240, 3)

    canvas.save(ARCH, quality=95)


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def margins(section, top=1.7, bottom=1.6, left=1.7, right=1.7):
    section.top_margin = Cm(top)
    section.bottom_margin = Cm(bottom)
    section.left_margin = Cm(left)
    section.right_margin = Cm(right)


def add_page_number(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    keep_with_next(p)
    return p


def add_para(doc, text="", bold_start=None, style=None):
    p = doc.add_paragraph(style=style)
    if bold_start and text.startswith(bold_start):
        p.add_run(bold_start).bold = True
        p.add_run(text[len(bold_start):])
    else:
        p.add_run(text)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.08
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)
        p.paragraph_format.space_after = Pt(3)


def add_picture_with_alt(run, path, width, alt):
    shape = run.add_picture(str(path), width=width)
    shape._inline.docPr.set("descr", alt)
    shape._inline.docPr.set("title", alt)
    return shape


def add_table(doc, headers, rows, widths=None, font_size=8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    prevent_row_split(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(BLACK)
        run.font.size = Pt(font_size)
        if widths:
            cell.width = Inches(widths[i])
    for row in rows:
        table_row = table.add_row()
        prevent_row_split(table_row)
        cells = table_row.cells
        for i, value in enumerate(row):
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(str(value))
            r.font.size = Pt(font_size)
            if widths:
                cells[i].width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


papers = [
    {
        "title": "IRIS: LLM-Assisted Static Analysis for Detecting Security Vulnerabilities",
        "venue": "ICLR 2025 (A*)",
        "data": "CWE-Bench-Java: 120 real Java vulnerabilities, 4 CWEs",
        "method": "LLM-inferred taint specifications + CodeQL + LLM contextual filtering",
        "metrics": "Detected vulnerabilities, FDR, F1, source/sink quality",
        "results": "GPT-4 version found 55/120 vs CodeQL 27/120; average FDR improved by about 5 points; four new vulnerabilities reported.",
        "limits": "Only Java and four taint-style CWEs; many misses; high FDR; several model calls and long runs; weak models can make filtering worse.",
        "summary": "IRIS tackles a common weakness of static analysis: the tool needs security specifications, but writing these specifications by hand is slow and incomplete. Its method is neuro-symbolic. An LLM first labels possible source and sink APIs, CodeQL then performs repository-level taint tracking, and another LLM checks the warnings with code context. The authors created CWE-Bench-Java with 120 manually checked vulnerabilities from real Java projects in four CWE groups. With GPT-4, IRIS detected 55 vulnerabilities, while the standard CodeQL setup detected 27. Its reported average false-discovery rate was still high, but about five percentage points better than CodeQL. A manual check of 50 reported alarms found 27 useful attack surfaces, and the system also found four previously unknown vulnerabilities. The paper is a strong base because it joins deterministic analysis with LLM judgment instead of trusting a model alone. Still, it only studies Java and taint-style problems, misses more than half of the benchmark, needs many LLM calls, and can take hours on some projects. Its filter is also not automatically safe: smaller models sometimes reduce precision. Our project keeps the hybrid idea but starts from existing SAST alerts, adds explicit evidence, abstention, and a CI budget instead of trying to discover every vulnerability from scratch."
    },
    {
        "title": "LLMxCPG: Context-Aware Vulnerability Detection Through Code Property Graph-Guided Large Language Models",
        "venue": "USENIX Security 2025 (A*)",
        "data": "FormAI-v2, PrimeVul, SVEN, ReposVul, and 60 post-cutoff CVEs",
        "method": "Joern CPG slicing + fine-tuned query generator and QwQ-32B classifier",
        "metrics": "Accuracy, precision, recall, F1, reduction ratio, robustness",
        "results": "Reduced code by 67.84–90.93% and improved F1 by 15–40% over reported baselines; ReposVul F1 0.610.",
        "limits": "Mostly memory-related C/C++ CWEs; A100 fine-tuning; per-dataset threshold tuning; weaker performance for rare CWEs and deep nesting.",
        "summary": "LLMxCPG tries to give an LLM the useful program context without sending an entire repository. Joern builds a code property graph, generated CPG queries find possible security paths, backward slicing removes unrelated code, and a fine-tuned model classifies the slice. The study uses FormAI-v2 and PrimeVul for training, then tests on SVEN, ReposVul, and 60 CVEs published after the model knowledge cutoff. Slices were 67.84% to 90.93% smaller than the original code and the paper reports F1 improvements of 15% to 40% over earlier systems. On ReposVul, the average accuracy was 0.634 and F1 was 0.610; the post-cutoff set reached F1 0.617. The robustness experiments are useful because they rename identifiers, insert statements, move code, and remove comments. However, the selected CWEs are mainly memory-safety problems in C/C++. The system needs Joern, generated queries, two 32B models and A100-class hardware. Thresholds were calibrated separately for each dataset, sometimes with only 20 labeled examples, so deployment calibration remains uncertain. Accuracy fell sharply at very deep nesting and results varied for rare CWEs. For our work, this paper supports targeted repository context and threshold calibration, but our agent uses cheaper read-only tools and measures the cost of each context request."
    },
    {
        "title": "PATCHAGENT: A Practical Program Repair Agent Mimicking Human Expertise",
        "venue": "USENIX Security 2025 (A*)",
        "data": "178 reproducible vulnerabilities from 30 C/C++ programs",
        "method": "LLM repair agent with clangd, patch verifier and four interaction optimizations",
        "metrics": "Valid repair rate, security/functional tests, ablations, tokens, cost, time",
        "results": "Union of five models repaired 164/178 (92.13%); best single model 84.83%; union averaged $6.11 per successful task.",
        "limits": "Requires a PoC and functional tests; expensive repeated runs; limited direct baseline comparison; plausible tests cannot prove semantic correctness.",
        "summary": "PATCHAGENT treats repair as a full agent task instead of only asking an LLM to write a patch. It performs fault localization, follows definitions using clangd, proposes edits, compiles them, replays a proof-of-concept input, and runs functional tests. Four extra controls clean noisy reports, compress the interaction history, correct malformed actions, and feed failed patches back as counterexamples. The evaluation contains 178 real vulnerabilities from OSS-Fuzz, Huntr and ExtractFix across nine low-level bug types. The best single model, Claude-3 Opus, repaired 84.83%, while taking the union across five models reached 92.13% or 164 cases. Ablation results show every interaction feature helped. The cost is important: the union averaged roughly 843k input tokens, 108k output tokens and $6.11 per successful case, and full ablation work was restricted because 75 cases cost more than $1,500. The setup also assumes a reproducible exploit and useful functional tests, which SAST warnings usually do not provide. Comparison with older repair tools was possible on only 28 compatible cases, and passing available tests does not prove a patch is correct for every input. It is relevant because it shows why tool feedback and state matter, while also motivating our strict budget and audit trail for CI triage."
    },
    {
        "title": "Benchmarking LLMs and LLM-based Agents in Practical Vulnerability Detection for Code Repositories",
        "venue": "ACL 2025 (A*)",
        "data": "JITVul: 1,758 paired C/C++ samples from 879 CVEs and 91 CWEs",
        "method": "Plain LLM, fixed dependency context, and ReAct agents with caller/callee tools",
        "metrics": "F1 and pairwise accuracy (pAcc)",
        "results": "ReAct improved pAcc over other methods by up to 16.61%, but absolute pAcc remained low (best reported 20.17%).",
        "limits": "Function target is given; C/C++ heavy; vul-intro tracing can be wrong; low absolute performance; agents depend on instruction following.",
        "summary": "This paper asks whether a detector can tell a vulnerable function from its patched version inside a real repository. The authors built JITVul with 879 vulnerable/benign pairs, 91 CWE types, and large C/C++ repositories. They compare a plain LLM, an LLM given a fixed set of dependencies, and a ReAct agent that can fetch callers and callees when needed. Standard F1 looked reasonable in some settings, but pairwise accuracy was much lower. The best reported pAcc was only 20.17%, which shows that many models label both versions as vulnerable. ReAct usually improved pAcc, by as much as 16.61 percentage points over another method, because it retrieved interprocedural context adaptively. The paper’s strongest point is its evaluation argument: F1 by itself can reward an unsafe prediction bias, so a paired or safety-focused metric is needed. Its limits are also serious. The target function is supplied, most repositories are C/C++, tracing the introducing commit is not always exact, and the absolute results are not ready for an automatic security gate. Smaller open models sometimes failed to follow the required agent format. For our project, JITVul supports adaptive tool use, but more importantly it justifies measuring missed real vulnerabilities, coverage and abstention instead of reporting only F1."
    },
    {
        "title": "Reducing False Positives in Static Bug Detection with LLMs: An Empirical Study in Industry",
        "venue": "ICSE 2026 SEIP (A* venue)",
        "data": "433 Tencent C/C++ alerts: 105 real, 328 false; three bug types",
        "method": "Prompts, few-shot variants, LLM4SA and LLM4PFA hybrid analysis",
        "metrics": "Accuracy, precision, recall, F1, FP reduction, time and API cost",
        "results": "Hybrid methods removed 94–98% of false alarms; best accuracy 0.94 and F1 0.86; recall stayed below the desired 0.90.",
        "limits": "Private one-company dataset; only three bug types; C/C++ focus; no release of data; randomness and long-context failures.",
        "summary": "Du et al. study the exact operational pain behind this project: developers spend around 10 to 20 minutes checking each static-analysis warning. They collected 433 alerts from Tencent, including 328 false positives and 105 real bugs, covering null dereference, out-of-bounds and divide-by-zero cases. Basic prompts, chain-of-thought and few-shot prompts were compared with hybrid methods that add static-analysis facts. The hybrid LLM4PFA method performed best, removing 94% to 98% of false alarms depending on the model. Its accuracy reached 0.93 to 0.94 and F1 reached 0.86, while processing one alert cost about 2.1 to 109.5 seconds and $0.0011 to $0.12. This is strong evidence that a hybrid approach can reduce review work. The safety concern is that recall stayed below Tencent’s desired 90%, meaning some real bugs would still be filtered. The data comes from one private business line and cannot be released, so replication and generalization are limited. The study covers only three C/C++ bug types and reports failures on long context, cascaded conditions and uncommon syntax. It strongly motivates our project, but we avoid using it as the base until the supervisor confirms that an ICSE industry track is acceptable."
    },
    {
        "title": "Fixing Security Vulnerabilities with Agentic AI in OSS-Fuzz",
        "venue": "ICSE 2026 SEIP (A* venue)",
        "data": "OSS-Fuzz historical benchmark plus 45 public unpatched vulnerabilities",
        "method": "CodeRover-S with dynamic call graph, type context and retry feedback",
        "metrics": "Plausible and manually correct patches, compilation, ablations",
        "results": "61.1% plausible with o3-mini, 71.8% with Gemini 2.5 Flash; only 45.5% of inspected plausible patches were judged correct.",
        "limits": "Exploit-based C/C++ setting; plausibility overestimates correctness; manual review subset; repair rather than SAST triage.",
        "summary": "This paper adapts AutoCodeRover to security bugs discovered by OSS-Fuzz. Instead of depending on a natural-language issue alone, CodeRover-S runs the exploit, extracts a dynamic call graph, summarizes relevant types, and carries failed locations and patches into later attempts. On the historical benchmark, the o3-mini version generated plausible patches for 61.1% of cases, compared with 39.0% for Agentless. Changing the model to Gemini 2.5 Flash raised plausibility to 71.8%. The authors also tested 45 publicly disclosed but still unpatched OSS-Fuzz issues; the system produced plausible patches for 73.3%, and five were merged by maintainers at the time of writing. The critical result is the manual correctness check: only 40 of 88 plausible patches, or 45.5%, were judged actually correct. Many merely stopped the supplied crashing input or changed unrelated logic. This shows why a convenient automatic metric can be misleading. The setting also assumes an executable exploit and mainly C/C++ projects, while our SAST triage has static alerts without a crash. Even so, the study gives practical support for adaptive context, feedback and human review. Our CI gate follows the same safety lesson by keeping uncertain alerts visible instead of treating a confident-looking output as proof."
    },
    {
        "title": "AutoCodeRover: Agentic Program Repair for SonarQube Issues",
        "venue": "FSE 2026 Industry (A* venue)",
        "data": "Industrial SonarQube deployment and reported user examples",
        "method": "Issue partitioning, isolated repair agents, SonarQube reruns and regression checks",
        "metrics": "Issues handled, time saved, deployment feedback and patch acceptance",
        "results": "A public example handled 26 issues in about 10 minutes versus an estimated 91 manual minutes; client feedback reported large debt reduction.",
        "limits": "No controlled public benchmark or full confusion matrix; confidentiality limits detail; repair success depends on SonarQube validation.",
        "summary": "AutoCodeRover for SonarQube describes an industrial system that turns static-analysis issues into repair tasks. It groups issues according to whether they can be fixed independently, launches isolated agents, uses repository navigation tools, applies patches in stages, and reruns SonarQube and regression checks before presenting changes. The design includes practical controls such as sandboxed tools, comment obfuscation to reduce prompt-injection risk, no persistent cross-task memory, and author approval before patches are accepted. The paper reports a public demonstration where 26 issues were handled in about ten minutes, compared with an estimated 91 minutes of manual work, and gives deployment feedback about reduced technical debt. It is very relevant to the target product because it sits directly after a static analyzer in a development workflow. However, the evidence is weaker than a normal benchmark study. Confidential client data prevents a reproducible dataset, there is no full precision/recall table, and the reported time savings depend on manual estimates and selected examples. SonarQube reruns can confirm that a rule no longer fires, but that alone cannot prove a semantic fix is safe. We therefore use this paper as industrial design evidence, not as the main base or the only support for effectiveness."
    },
    {
        "title": "Let the Trial Begin: A Mock-Court Approach to Vulnerability Detection using LLM-Based Agents",
        "venue": "ICSE 2026 Research (A*)",
        "data": "PrimeVul paired functions; 435 evaluation pairs; limited tuning data",
        "method": "Four roles: security researcher, code author, moderator and review board",
        "metrics": "Pairwise correct/error categories, precision, recall and FPR",
        "results": "GPT-4o VulTrial reached 81 correct pairs vs 40 for single-agent CoT and 44 for GPTLens; FPR was 0.52.",
        "limits": "High false-positive rate; function-level input; multiple calls; same base model for all roles; small role-tuning set.",
        "summary": "VulTrial uses a courtroom-style debate to reduce one-sided model reasoning. A security researcher argues that the function is vulnerable, a code author gives a defense, a moderator summarizes the dispute, and a review board returns the verdict. The evaluation uses vulnerable/benign pairs from PrimeVul, where the two functions differ only slightly. With GPT-4o, VulTrial classified 81 of 435 pairs fully correctly, compared with 40 for a single-agent chain-of-thought prompt and 44 for GPTLens. Tuning individual roles with only 50 validation pairs improved the pairwise score further, and an ablation study showed that removing roles hurt performance. The paper also rates explanation quality and reports confirmed zero-day findings. The improvement is meaningful, but the absolute result still means most pairs were not both correct. The GPT-4o version had a false-positive rate of 0.52, used several model calls per function, and only received function-level code rather than a full CI alert and repository state. All roles normally use the same underlying model, so the debate may repeat the same bias rather than provide truly independent evidence. For our project, it suggests that challenge and counterargument can help, but a full four-agent court is too costly. We instead use tools plus evidence and reserve human review for uncertainty."
    },
    {
        "title": "Out of Distribution, Out of Luck: How Well Can LLMs Trained on Vulnerability Datasets Detect Top 25 CWE Weaknesses?",
        "venue": "ICSE 2026 Research (A*)",
        "data": "TitanVul (38,548 functions) and BenchVul (1,050 vulnerable/fix pairs)",
        "method": "Dataset audit, multi-agent validation, OOD benchmark and realistic synthesis",
        "metrics": "ID/OOD accuracy, CWE coverage and manual validity",
        "results": "BigVul-trained Qwen scored 0.703 ID but 0.493 OOD real; TitanVul scored 0.590 ID but 0.881 OOD, rising to 0.932 with synthesis.",
        "limits": "Function-level task; BenchVul is 81.9% synthetic; manual correctness 92%; training/evaluation may not match SAST alerts.",
        "summary": "This study shows that a high score on a familiar vulnerability dataset can give a false sense of security. The authors audit seven public datasets, then build TitanVul with 38,548 validated vulnerable functions and BenchVul with 1,050 vulnerable/fix pairs across the MITRE Top 25. A multi-agent pipeline checks training pairs, while a separate four-role workflow generates missing examples for rare CWEs. The main result is the gap between in-distribution and independent testing. Qwen trained on BigVul scored 0.703 on its own distribution but only 0.493 on BenchVul’s real subset, close to chance. Training on TitanVul gave a lower in-distribution score of 0.590 but much better out-of-distribution accuracy of 0.881; synthetic augmentation raised this to 0.932. This is strong evidence that dataset design matters more than one headline metric. The benchmark is still function-level, its manual correctness estimate is 92%, and 81.9% of BenchVul’s examples are synthetic because real samples were scarce. A function classifier also differs from triaging concrete SAST paths. For our work, the paper warns us not to claim generality from OWASP Benchmark alone. We will report per-CWE results, use a fixed held-out set, and present external validation as future work rather than hiding the limitation."
    },
    {
        "title": "APPATCH: Automated Adaptive Prompting Large Language Models for Real-World Software Vulnerability Patching",
        "venue": "USENIX Security 2025 (A*)",
        "data": "306 training examples, 97 zero-day vulnerabilities, and ExtractFix",
        "method": "Dynamic exemplar selection, semantic slicing, multi-faceted validation",
        "metrics": "Recall, precision and F1 for syntactic, semantic, plausible and correct patches",
        "results": "Best zero-day correctness F1 was 36.46%; best ExtractFix correctness F1 was 73.86%; slicing and validation helped.",
        "limits": "Assumes CWE/location information; six C CWEs; small manually labeled sets; upstream CodeQL errors lower end-to-end performance.",
        "summary": "APPATCH improves vulnerability repair prompts by selecting relevant examples at run time and reducing the input to a security-focused slice. It first reasons about vulnerability semantics, retrieves similar fixes, asks an LLM for several candidates, and filters them with multiple validation checks. The authors use 306 labeled examples to support prompting and evaluate on 97 vulnerabilities disclosed after the model cutoff plus the smaller ExtractFix set. Across GPT-4, Gemini 1.5, Claude 3.5 and Llama 3.1, APPATCH usually beats direct reasoning, fixed examples and zero-shot completion. The best correctness F1 on the difficult zero-day set was 36.46% with Claude 3.5, while the same model reached 73.86% on ExtractFix. Removing semantic slicing or using random examples reduced performance. The gap between the two datasets is an honest sign that new real vulnerabilities remain hard. APPATCH assumes the CWE and bug location are already known; when CodeQL supplies them, upstream mistakes reduce end-to-end accuracy. It covers six C-language CWEs, uses small manually labeled datasets, and many cases are not fully compilable with tests. This paper supports focused context and independent validation. Our project applies the same principle before the gate: retrieve only relevant repository evidence, but do not suppress an alert when the evidence is incomplete."
    },
    {
        "title": "Revisiting Security Practices for GitHub Actions Workflows",
        "venue": "ICPC 2025 ERA (A)",
        "data": "18,938 workflows from 5,246 active public repositories, 2022 vs 2024",
        "method": "Enhanced GHAST scan of two workflow snapshots",
        "metrics": "Issue counts, affected workflows, exploitability and Wilcoxon tests",
        "results": "33,362 issues in 2024; 74.9% omitted explicit permissions; confidential-data findings increased from 2,345 to 2,617.",
        "limits": "Four rule types, public repositories only, scanner validity checked on 40 reports, configuration issues rather than code SAST alerts.",
        "summary": "Huang and Lin examine the security of the CI/CD environment itself. They collected 18,938 GitHub Actions workflows from 5,246 active public repositories and compared 2022 with 2024. An enhanced GHAST scanner checked confidential-data exposure, command injection, unpinned unverified actions, and missing token permissions. It found 33,362 issues in the 2024 workflows. Permission practice improved: cases with no explicit permissions fell from 16,036 to 14,192. Even then, 74.9% of workflows still omitted a permission declaration, while only 9.5% used job-level permissions. Confidential-data findings increased from 2,345 to 2,617 and command-injection findings rose slightly, with no statistically significant improvement for vulnerabilities. The paper is only five pages and covers four rule families. Its scanner was manually checked on 40 reports, which is too small to rule out false negatives, and the sample excludes private repositories. It also studies workflow configuration rather than Java SAST alerts. Still, it is directly important to our product: an accurate agent is not enough if its GitHub Actions workflow has broad tokens or unsafe third-party actions. We therefore include least-privilege permissions, pinned dependencies where possible, read-only analysis tools, and no automatic code write access in the system design."
    },
]


references = [
    "[1] Z. Li, S. Dutta, and M. Naik, “IRIS: LLM-Assisted Static Analysis for Detecting Security Vulnerabilities,” in Proc. ICLR, 2025. https://proceedings.iclr.cc/paper_files/paper/2025/hash/582d4e27fa24168f3af1f4582655034b-Abstract-Conference.html",
    "[2] A. Lekssays, H. Mouhcine, K. Tran, T. Yu, and I. Khalil, “LLMxCPG: Context-Aware Vulnerability Detection Through Code Property Graph-Guided Large Language Models,” in Proc. 34th USENIX Security Symposium, 2025, pp. 489–506.",
    "[3] Z. Yu et al., “PATCHAGENT: A Practical Program Repair Agent Mimicking Human Expertise,” in Proc. 34th USENIX Security Symposium, 2025, pp. 4381–4398.",
    "[4] A. Yildiz, S. G. Teo, Y. Lou, Y. Feng, C. Wang, and D. M. Divakaran, “Benchmarking LLMs and LLM-based Agents in Practical Vulnerability Detection for Code Repositories,” in Proc. ACL, 2025, pp. 30848–30865.",
    "[5] X. Du et al., “Reducing False Positives in Static Bug Detection with LLMs: An Empirical Study in Industry,” in Proc. ICSE-SEIP, 2026.",
    "[6] Y. Zhang, J. Wang, D. Berzin, M. Mirchev, and A. Roychoudhury, “Fixing Security Vulnerabilities with Agentic AI in OSS-Fuzz,” in Proc. ICSE-SEIP, 2026, doi: 10.1145/3786583.3786880.",
    "[7] M. Mirchev, R. Shariffdeen, H. Ruan, Y. Zhang, and A. Roychoudhury, “AutoCodeRover: Agentic Program Repair for SonarQube Issues,” in Proc. FSE Companion, 2026, doi: 10.1145/3803437.3805209.",
    "[8] R. Widyasari et al., “Let the Trial Begin: A Mock-Court Approach to Vulnerability Detection using LLM-Based Agents,” in Proc. ICSE, 2026, doi: 10.1145/3744916.3773256.",
    "[9] Y. Li et al., “Out of Distribution, Out of Luck: How Well Can LLMs Trained on Vulnerability Datasets Detect Top 25 CWE Weaknesses?” in Proc. ICSE, 2026.",
    "[10] Y. Nong, H. Yang, L. Cheng, H. Hu, and H. Cai, “APPATCH: Automated Adaptive Prompting Large Language Models for Real-World Software Vulnerability Patching,” in Proc. 34th USENIX Security Symposium, 2025, pp. 4479–4497.",
    "[11] J. Huang and B. Lin, “Revisiting Security Practices for GitHub Actions Workflows,” in Proc. 33rd IEEE/ACM ICPC, 2025, pp. 73–77, doi: 10.1109/ICPC66645.2025.00016.",
    "[12] OWASP Foundation, “OWASP Benchmark,” version 1.2. https://owasp.org/www-project-benchmark/",
    "[13] ICORE Conference Portal, “Conference Rankings,” 2026. https://portal.core.edu.au/conf-ranks/",
]


def build():
    make_architecture()
    doc = Document()
    for sec in doc.sections:
        margins(sec)
        add_page_number(sec)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    for name, size in [("Title", 24), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)]:
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(4)
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    styles["List Bullet"].font.name = "Times New Roman"
    styles["List Bullet"].font.size = Pt(11)

    # Cover
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(95)
    p_pr = p._p.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)
    r = p.add_run("GateKeeper AI")
    r.bold = True
    r.font.size = Pt(26)
    r.font.color.rgb = RGBColor.from_string(BLACK)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run("Evidence-Grounded Agentic Triage for Safer SAST Gates in CI/CD")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor.from_string(BLACK)
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(30)
    p3.add_run("AI4015 Agentic AI - Fall 2026\nSemester Project - Phase 1\n\nSelected seed: G3 - Empowering CI/CD with Agentic AI for Enhanced Software Security\nGroup identifier: G3-10").font.size = Pt(12)
    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.space_before = Pt(55)
    rr = p4.add_run("Group Members\n\nAnas Imran - 22I-1319 (Group Head)\nHamza Farooq - 22I-0912\nMaaz Bhatti - 22I-1378")
    rr.font.size = Pt(12)
    doc.add_page_break()

    add_heading(doc, "Executive summary", 1)
    add_para(doc, "Static-analysis tools are useful in CI/CD because they can find security problems before code is merged. The problem is that a noisy scanner can also block safe changes. In our OWASP Benchmark run, FindSecBugs v1.4.6 produced 2,113 alerts. About 35% were false alarms, even though average true-positive rate was 0.968. This creates a practical choice between wasting developer time and ignoring warnings.")
    add_para(doc, "GateKeeper AI is a safety-first triage layer between SAST and the merge gate. It compares three systems: B0 keeps every alert, B1 asks an LLM once with no tools, and the proposed agent can inspect repository context with read-only tools. The agent records evidence, abstains when confidence is low, and stops under a fixed budget. A true-positive verdict blocks the merge, a false-positive verdict suppresses the alert with a reason, and an unsure verdict goes to a human.")
    add_para(doc, "The literature review covers eleven peer-reviewed papers published in 2025 or 2026 at ICLR, USENIX Security, ACL, ICSE, FSE and ICPC. IRIS [1] is the base paper because it is an ICLR 2025 paper at an A* venue and gives a usable hybrid static-analysis comparison. The three selected gaps are evidence-grounded triage, calibrated abstention for safe CI decisions, and resource-aware agent control. The technical foundation includes a real-data pipeline, two model-facing baselines, the agent loop, CI gate, structured traces, evaluation metrics and eleven passing tests.")

    add_heading(doc, "Contents", 1)
    add_bullets(doc, [
        "1. Project and paper selection",
        "2. Base paper study",
        "3. Critical literature review",
        "4. Cross-paper comparison",
        "5. Research-gap synthesis",
        "6. Complete project proposal",
        "7. Technical foundation and initial results",
        "8. Risks, milestones, work division and publication direction",
        "9. References",
        "Appendix A: ranking evidence",
    ])

    add_heading(doc, "1. Project and paper selection", 1)
    add_heading(doc, "1.1 Project title and seed", 2)
    add_para(doc, "Selected seed: G3 — Empowering CI/CD with Agentic AI for Enhanced Software Security.")
    add_para(doc, "Proposed title: GateKeeper AI: Evidence-Grounded Agentic Triage for Safer SAST Gates in CI/CD.")
    add_heading(doc, "1.2 Selection audit", 2)
    add_para(doc, "All eleven counted papers are peer-reviewed conference papers from 2025 or 2026. No arXiv-only preprint is counted. Some 2026 papers were first circulated as author manuscripts, but their official conference pages show acceptance in the named tracks. ICORE/CORE evidence is included in Appendix A. The rank applies to the conference series, while the paper list states the research, industry, SEIP or ERA track separately.")
    audit_rows = [[str(i + 1), p["title"], p["venue"], "Yes" if i == 0 else "No"] for i, p in enumerate(papers)]
    add_table(doc, ["#", "Paper", "Year / venue / rank", "Base"], audit_rows, [0.35, 4.25, 1.65, 0.45], 7.5)
    add_para(doc, "The 2026 Empirical Software Engineering article “Cracking IoT security: can LLMs outsmart static analysis tools?” was also verified through Springer (doi: 10.1007/s10664-026-10849-9). It was not counted because its IoT focus is less direct than the selected CI/CD and repository-security studies, and the selected list already exceeds ten without needing a separate journal-quartile proof.")

    add_heading(doc, "2. Base paper study: IRIS", 1)
    add_heading(doc, "2.1 Full reference and rank", 2)
    add_para(doc, references[0])
    add_para(doc, "ICLR is ranked A* in the 2026 ICORE conference portal. The captured official evidence appears in Appendix A. The full 24-page proceedings PDF was read for this report, including the appendices and detailed failure analysis.")
    add_heading(doc, "2.2 Why IRIS is the base paper", 2)
    add_para(doc, "IRIS is a better base than the ICSE industry paper because there is no track-eligibility doubt and its artifacts are public. Its method is also close enough to compare against: both IRIS and GateKeeper AI combine symbolic analysis with LLM judgment, use repository context, and aim to reduce the manual effort caused by static analysis. The difference is the decision point. IRIS builds new taint specifications and discovers flows; our project starts from SAST warnings already produced in CI and decides whether each warning should block, be suppressed, or be reviewed.")
    add_heading(doc, "2.3 Problem, architecture and benchmark", 2)
    add_para(doc, "IRIS addresses the shortage of hand-written source, sink and taint specifications. First, it extracts external API calls and internal method parameters that could be security-relevant. An LLM labels them for a selected CWE. CodeQL uses these inferred specifications to compute source-to-sink flows across the repository. A final contextual-analysis prompt then inspects each flow and tries to remove false alarms. This division is sensible: the LLM supplies flexible domain knowledge, while CodeQL supplies repeatable graph reasoning.")
    add_para(doc, "The benchmark, CWE-Bench-Java, contains 120 manually validated vulnerabilities in 120 compilable Java projects: 55 path traversal cases (CWE-22), 13 command injection cases (CWE-78), 31 cross-site scripting cases (CWE-79), and 21 code injection cases (CWE-94). Projects average roughly 300 thousand lines of code and ten exceed one million lines. Baselines include CodeQL 2.15.3, Infer, SpotBugs and Snyk, and the study compares seven LLMs.")
    add_heading(doc, "2.4 Main results", 2)
    add_bullets(doc, [
        "IRIS with GPT-4 detected 55 of 120 known vulnerabilities; CodeQL detected 27.",
        "Detection rate increased from 22.5% to 45.83%. Average F1 increased from 0.076 to 0.177.",
        "Average reported false-discovery rate was 84.82% for GPT-4 IRIS versus 90.03% for CodeQL. This is an improvement, but still too noisy for an automatic blocking gate.",
        "In a manual sample of 50 IRIS alarms, 27 were judged potential attack surfaces, suggesting an estimated FDR of about 46% because the benchmark labels only the known vulnerability.",
        "The authors found four previously unknown vulnerabilities in the latest versions of 30 projects.",
    ])
    add_heading(doc, "2.5 Limitations and how this project goes beyond it", 2)
    add_para(doc, "The authors openly state that IRIS still misses many vulnerabilities, makes many LLM calls, can be expensive, and is only evaluated on Java. The detailed errors show more limits. Some bugs are not simple taint flows. Side effects and propagators may be missing from the graph. Candidate extraction excludes several return, exception and static-call patterns. Very large projects can take more than an hour, and a few take several hours. Contextual filtering only helps when the model is capable enough; weaker models can remove useful results or add noise.")
    add_para(doc, "Our group adds three controls that IRIS does not evaluate as a CI decision system. First, every suppressed alert must carry file-level evidence gathered by read-only tools. Second, low-confidence decisions become UNSURE and stay visible to a human. Third, each episode has a tool and model budget that is measured together with latency and tokens. We also compare against a no-tool single call, report confidence calibration, and treat real-vulnerability retention as a hard safety measure.")

    add_heading(doc, "3. Critical summaries of the reviewed papers", 1)
    for idx, paper in enumerate(papers, 1):
        add_heading(doc, f"3.{idx} {paper['title']}", 2)
        add_para(doc, paper["summary"])

    add_heading(doc, "3.12 Cross-paper synthesis", 2)
    add_para(doc, "Three patterns repeat across the review. First, useful systems combine an LLM with a source of program facts. IRIS uses CodeQL, LLMxCPG uses code property graphs, PATCHAGENT uses a language server and tests, and the OSS-Fuzz agent uses execution and type information. The plain-prompt baselines are generally weaker or less dependable. This supports our decision to let the model request repository evidence instead of treating the alert text as the full problem.")
    add_para(doc, "Second, the evaluation metric can change the conclusion. JITVul shows that an apparently acceptable F1 score can come from labeling nearly everything vulnerable. The OSS-Fuzz study shows that a patch can stop one crashing input but still be wrong. TitanVul shows that high in-distribution accuracy can collapse on an independent benchmark. For a CI gate, the closest equivalent failure is a high suppression rate that hides real vulnerabilities. Our results will therefore place the count of wrongly suppressed real alerts beside every aggregate score.")
    add_para(doc, "Third, more agent activity is not automatically better. Debate, retries, larger context and model ensembles can improve results, but they also increase token use, latency and cost. GateKeeper AI will compare fixed budgets and early stopping so that any accuracy gain can be reported together with the extra work required.")

    # Landscape comparison table
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    margins(sec, 1.2, 1.2, 1.2, 1.2)
    add_heading(doc, "4. Cross-paper comparison table", 1)
    comp_rows = []
    for i, p in enumerate(papers, 1):
        title = ("BASE — " if i == 1 else "") + p["title"]
        comp_rows.append([i, title, p["venue"], p["data"], p["method"], p["metrics"], p["results"], p["limits"]])
    add_table(doc, ["#", "Title", "Year / venue", "Dataset / benchmark", "Method / agent", "Metrics", "Key results", "Limitations / gap"], comp_rows, None, 6.2)
    add_para(doc, "Critical reading across the table: the most impressive headline scores often use inputs that a normal SAST gate does not have, such as a crashing exploit, the exact vulnerable function, or the correct CWE and location. Several studies also show that plausible outputs, F1, or in-distribution accuracy can overstate practical safety. This is why our evaluation separates suppression from real-vulnerability retention and reports abstention, calibration, evidence and cost.")

    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    margins(sec)

    add_heading(doc, "5. Research-gap synthesis", 1)
    gap_rows = [
        ["G1. Evidence-grounded SAST triage", "IRIS still reports high FDR [1]; JITVul agents remain inaccurate [4]; Du et al. show static facts help [5].", "A read-only agent gathers exact code, definitions and nearby uses, then attaches evidence to every verdict.", "False-alarm suppression and real-vulnerability retention; evidence coverage; blinded manual validity of evidence."],
        ["G2. Calibrated abstention for CI", "JITVul shows F1 can hide unsafe prediction bias [4]; Du et al. remain below the desired recall [5]; TitanVul shows OOD confidence risk [9].", "Calibrate confidence on a validation split. Low-confidence cases become UNSURE and are never silently suppressed.", "ECE and Brier score; risk–coverage curve; missed real vulnerabilities; human-review load."],
        ["G3. Resource-aware agent control", "IRIS can require many calls and long runtimes [1]; PATCHAGENT is costly [3]; VulTrial uses several agents [8].", "Track steps, tokens and latency; impose tool budgets; use early stopping when enough evidence is collected.", "Accuracy/safety at budgets 4, 8 and 12; tokens, calls, tool steps, latency and cost per alert."],
    ]
    add_table(doc, ["Gap", "Supporting evidence", "Proposed approach", "How it will be measured"], gap_rows, [1.25, 2.05, 2.05, 1.65], 7.3)
    add_heading(doc, "5.1 Research questions, hypotheses and success criteria", 2)
    rq_rows = [
        ["RQ1", "Does adaptive repository evidence improve false-positive suppression over a no-tool LLM call without reducing safety?", "H1: the agent will improve suppression by at least 15 percentage points over B1 while retaining at least 95% of real vulnerabilities.", "Retention ≥0.95; suppression gain ≥0.15; valid evidence on ≥0.90 of decided cases."],
        ["RQ2", "Can calibrated abstention reduce unsafe suppressions while keeping human review manageable?", "H2: a tuned threshold will keep missed real vulnerabilities at or below 5% with review load at or below 35% and ECE below 0.10.", "Miss rate ≤0.05; review load ≤0.35; ECE <0.10 on held-out data."],
        ["RQ3", "How much agent budget is needed before added tool use stops improving decisions?", "H3: budget 8 will stay within two F1 points of budget 12 while using at least 25% fewer tool steps or tokens.", "ΔF1 ≤0.02 vs budget 12; ≥25% efficiency saving; median latency within the CI target."],
    ]
    add_table(doc, ["", "Research question", "Hypothesis", "Success criterion"], rq_rows, [0.45, 2.55, 2.45, 1.55], 7.3)

    add_heading(doc, "6. Complete project proposal", 1)
    add_heading(doc, "6.1 Problem statement and motivation", 2)
    add_para(doc, "The target users are developers and application-security reviewers who receive SAST warnings on pull requests. A conventional gate treats every warning as equally reliable. In our real baseline, this would keep 2,113 alerts active and about 740 of them are false. Developers then either spend time reviewing noise or learn to ignore the tool. The security risk is not only inconvenience: a careless filter can solve alert fatigue by hiding a real vulnerability. The project therefore treats triage as a selective decision problem with unequal costs. Suppressing a false alert is useful; suppressing a real vulnerability is unacceptable.")
    add_heading(doc, "6.2 Target product", 2)
    add_para(doc, "The semester product will be a command-line service and GitHub Actions integration that ingests SpotBugs/FindSecBugs XML, normalises findings, gathers repository evidence, produces TP/FP/UNSURE verdicts, and applies a safe CI policy. It will save a JSONL audit record for each alert, including rationale, evidence, tool steps, confidence, tokens and latency. The final dashboard is the GitHub job summary and downloadable metrics, not a separate web application.")
    add_para(doc, "Repository: https://github.com/Cheemaboi/gatekeeper-ai-sast-triage")
    add_heading(doc, "6.3 Why agentic behavior is needed", 2)
    add_para(doc, "A deterministic workflow can parse alerts and enforce a gate, but it cannot decide which surrounding code matters for every warning. A single LLM call has the opposite problem: it can reason, but it only sees the fixed context chosen in advance. A simple RAG system retrieves similar text, while vulnerability evidence often depends on exact definitions, callers, sanitizers and control flow in the current repository. The proposed agent can decide whether to read the flagged file, search for a symbol, inspect a definition, or stop. It keeps state across observations, works under a budget, and abstains when the available evidence does not justify suppression. These are goal-directed choices based on intermediate results, which is the part that genuinely needs agentic behavior.")
    add_heading(doc, "6.4 Goals", 2)
    add_bullets(doc, [
        "Reduce the number of false SAST alerts reaching developers while preserving at least 95% of known real vulnerabilities.",
        "Make every automated suppression reviewable through concrete repository evidence and a short rationale.",
        "Calibrate confidence so uncertainty increases human review instead of increasing hidden risk.",
        "Keep average resource use suitable for a pull-request job and expose cost, latency and tool efficiency.",
        "Deliver a public, reproducible repository with documented setup and evaluation commands.",
    ])
    add_heading(doc, "6.5 Confirmed data and environment", 2)
    add_para(doc, "The main environment is OWASP Benchmark v1.2 for Java. The pipeline reads its expected-results CSV, Java source tree, and FindSecBugs v1.4.6 XML. The full benchmark contains 2,740 labeled test cases and the scanner produced 2,113 deduplicated alerts. A fixed, seed-42 evaluation file contains 209 alerts: 121 real vulnerabilities and 88 false alarms, stratified by category and label. Initial model runs use the first 50 records from this fixed file. The repository itself is the agent’s read-only tool environment. No dataset will be selected later.")
    add_heading(doc, "6.6 System architecture and controls", 2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_picture_with_alt(
        p.add_run(), ARCH, Inches(6.75),
        "GateKeeper AI architecture showing B0 raw SAST, B1 single-call LLM, and the proposed agent. The agent marks G1 at evidence tools, G2 at confidence-based abstention, and G3 at budgeted state."
    )
    cap = doc.add_paragraph("Figure 1. GateKeeper AI architecture. G1, G2 and G3 mark the selected research gaps; B0 and B1 show the comparison paths.")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(8)
    add_bullets(doc, [
        "Models: planned local Qwen2.5-Coder 7B for the first reproducible run; OpenAI-compatible and Anthropic providers remain supported for later comparison.",
        "Tools: read_file, grep and find_definition, all restricted to the checked-out repository. There is no shell, network, write, commit or deployment tool in the triage loop.",
        "State: alert contract, conversation, observations, evidence list, cost units, tool steps, tokens, confidence and final verdict.",
        "Human control: UNSURE always enters review. FP suppression keeps its evidence. Repository code is not changed automatically.",
        "CI security: the example workflow should use least-privilege token permissions and pinned third-party actions where possible, following the risks in [11].",
    ])
    add_heading(doc, "6.7 Baselines and comparisons", 2)
    add_table(doc, ["System", "Information and behavior", "Purpose"], [
        ["B0: raw SAST", "Keeps every FindSecBugs alert; no LLM and no tools.", "Measures original alert burden and guarantees no triage-caused misses."],
        ["B1: single call", "One model call with alert metadata and flagged file; no adaptive tools.", "Separates the value of an LLM from the value of agentic context gathering."],
        ["Proposed agent", "Adaptive read-only tools, evidence, state, abstention and fixed budget.", "Tests G1–G3 and supplies the deployable CI gate."],
        ["Base-paper comparison", "Compare findings to IRIS’s hybrid LLM/static-analysis results and discuss task differences.", "Shows how the project builds on and goes beyond the base paper without claiming a direct leaderboard result."],
    ], [1.2, 3.6, 2.2], 8)
    add_heading(doc, "6.8 Evaluation plan", 2)
    add_bullets(doc, [
        "Primary evaluation: run B1 and the proposed agent on the same 209-alert file. Use a fixed initial subset of 50 only for Phase 1 evidence, then run the full set for the final project.",
        "Quality metrics: precision, decided recall, F1 and decided accuracy. Report per-CWE/category results because aggregate scores can hide injection-specific failures.",
        "Safety metrics: real-vulnerability retention, count of wrongly suppressed real alerts, false-alarm suppression, coverage, selective error and human-review load.",
        "Calibration metrics: ten-bin expected calibration error, Brier score and a risk–coverage curve as the confidence threshold changes.",
        "Evidence metrics: percentage of decided verdicts with evidence, average evidence items, and blinded manual checking of whether evidence actually supports the verdict.",
        "System metrics: LLM calls, tool steps, tokens, median latency and estimated cost per alert.",
        "Ablations: no tools (B1); no abstention; budgets 4/8/12; evidence requirement on/off; and, if time permits, one local and one hosted model.",
        "Analysis: use paired bootstrap confidence intervals for metric differences. Report raw counts beside percentages. Treat any decrease in real-vulnerability retention as a safety finding, even if F1 improves.",
    ])
    add_heading(doc, "6.9 Reproducibility and ethics", 2)
    add_para(doc, "OWASP Benchmark is downloaded from its original project and not redistributed. The evaluation file contains alert metadata and labels rather than benchmark source. Model temperature is zero where the provider supports it, every prompt and output is logged, and seeds are fixed. The tool boundary is read-only. No vulnerability is tested against a live third-party system, no exploit is generated, and no model decision automatically changes code. Provider data-handling terms must be checked before proprietary code is used; the local-model path is the privacy fallback.")

    add_heading(doc, "7. Technical foundation and initial results", 1)
    add_heading(doc, "7.1 Implemented components", 2)
    add_bullets(doc, [
        "Data loader for OWASP ground truth and SpotBugs/FindSecBugs XML, with deduplication and CWE/category normalization.",
        "Pydantic contracts for Finding, Evidence, Verdict, TriageRecord and GateDecision.",
        "B0 raw-SAST scorecard, B1 one-call LLM baseline, and a budgeted ReAct-style agent.",
        "Sandboxed read-only repository tools, confidence-based abstention and a safe CI gate.",
        "Evaluation metrics covering detection, triage safety, calibration, evidence, cost and latency.",
        "Eleven offline tests covering parsing, labeling, tools, both model loops, gate behavior and calibration. All eleven pass on 17 October 2026.",
    ])
    add_heading(doc, "7.2 B0 result on real data", 2)
    add_table(doc, ["Scanner", "Alerts", "Alert precision", "Avg TPR", "Avg FPR", "Youden"], [["FindSecBugs 1.4.6", "2,113", "0.648", "0.968", "0.609", "0.360"]], [1.6, 1.0, 1.2, 1.0, 1.0, 0.9], 8.5)
    add_para(doc, "The baseline finds nearly all known vulnerable cases on average, but 35.2% of its alerts are false. The worst categories are injection-related: XPath injection precision is 0.44, LDAP injection 0.50, command injection 0.53, path traversal 0.52 and SQL injection 0.56. This is exactly the alert mix where unsafe suppression would be tempting, so real-vulnerability retention remains the main constraint.")
    add_heading(doc, "7.3 Model integration", 2)
    add_para(doc, "The selected Phase 1 model is Gemini 3.5 Flash Lite through Google's OpenAI-compatible endpoint. It is used for the first 50-alert comparison. The implementation also supports local Ollama, OpenRouter and Anthropic providers. Each run produces JSONL verdict records, agent traces and a separate metrics file so that B1 and the proposed agent can be compared on the same evaluation sample.")
    add_para(doc, "The 50-alert B1 pilot produced 34 parseable decisions: precision 0.864, decided recall 0.905, F1 0.884, false-alarm suppression 0.769 and median latency 5.12 seconds. Sixteen cases were retained as errors because the free endpoint returned malformed output or a rate-limit response. The Gemini agent pilot completed all 50 records, producing 8 model decisions and retaining 42 cases as UNSURE with HTTP 429 errors after the free project quota was reached. The JSONL trace records both outcomes and supports a later higher-quota rerun.")

    add_heading(doc, "8. Milestones, feasibility, work division and publication", 1)
    add_heading(doc, "8.1 Twelve-week milestone plan", 2)
    add_table(doc, ["Weeks", "Deliverable", "Exit test"], [
        ["1–2", "Confirm paper set, base paper, gaps, dataset and contracts", "Supervisor accepts base/track interpretation; report sources frozen"],
        ["3–4", "Reproduce B0 and run 50-alert B1/agent pilot", "Both runs have complete JSONL, metrics and no parser failures"],
        ["5–6", "Evidence scoring and confidence calibration", "Manual evidence rubric agreed; threshold selected without using test labels"],
        ["7–8", "Full 209-alert comparison and per-CWE analysis", "B0/B1/agent tables plus paired bootstrap intervals"],
        ["9", "Ablations for budget, tools and abstention", "Budgets 4/8/12 and no-abstention results recorded"],
        ["10", "GitHub Actions integration and security hardening", "Least-privilege workflow runs on a pull request"],
        ["11", "Failure analysis and developer review", "All unsafe suppressions manually categorized"],
        ["12", "Final report, demo and artifact release", "Fresh clone reproduces tests and documented evaluation"],
    ], [0.7, 3.6, 2.7], 8)
    add_heading(doc, "8.2 Technical risks and fallback", 2)
    add_table(doc, ["Risk", "Effect", "Mitigation / fallback"], [
        ["Local model is too slow or does not fit memory", "Pilot cannot finish in CI time", "Use a smaller quantization, lower the tool budget, or run gpt-4o-mini and report hosted cost."],
        ["Model returns invalid JSON or tool calls", "Missing verdicts", "Strict parsing, one repair attempt, then safe UNSURE; count errors explicitly."],
        ["Agent suppresses a real vulnerability", "Security regression", "Tune for retention, require evidence, abstain, and fail safe to TP when abstention is disabled."],
        ["OWASP results do not generalize", "Weak external validity", "Report per-category results and add a small public real-project sample if time permits."],
        ["Tool context grows too large", "Cost and latency increase", "Bound file windows and grep results; compare budgets and early stopping."],
        ["GitHub workflow has broad permissions", "Supply-chain risk", "Set explicit read-only permissions and review action pinning, following [11]."],
    ], [1.55, 1.75, 3.7], 8)
    doc.add_page_break()
    add_heading(doc, "8.3 Work division", 2)
    add_para(doc, "The work is divided by primary responsibility, while all three members remain responsible for understanding the full project and reviewing the final results.")
    add_table(doc, ["Member", "Primary responsibilities", "Shared responsibility"], [
        ["Anas Imran (22I-1319), Group Head", "Paper selection, base-paper argument, supervisor coordination, final integration and submission.", "Review all results and sign off on claims."],
        ["Hamza Farooq (22I-0912)", "Provider setup, B1 and agent runs, calibration, ablations, metrics and failure analysis.", "Explain the architecture and reproduce B0."],
        ["Maaz Bhatti (22I-1378)", "CI workflow, security hardening, diagrams, ranking evidence, report formatting and reproducibility checks.", "Read the three gap papers and verify the tables."],
    ], [1.55, 3.45, 2.0], 8)
    add_heading(doc, "8.4 Potential publication direction", 2)
    add_para(doc, "If the results are strong, the publishable contribution is not another claim that an LLM can classify vulnerable code. It is a selective, evidence-grounded CI decision policy evaluated under asymmetric security cost. A useful paper would release the 209-alert triage benchmark, trace schema, evidence rubric and budget/calibration study, then show when an agent should act, abstain or stop. The natural first targets are a software-engineering security workshop, an ICSE/FSE tool-demo or NIER track, and later a full empirical paper after adding multiple projects, languages and developer feedback.")

    add_heading(doc, "9. References", 1)
    for ref in references:
        p = doc.add_paragraph(ref)
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        p.paragraph_format.space_after = Pt(4)
        p.runs[0].font.size = Pt(8.5)

    add_heading(doc, "Appendix A. ICORE/CORE ranking evidence", 1)
    evidence = [
        ("ICLR — ICORE2026 A*", "ICLR_CORE_Astar.png"),
        ("USENIX Security — ICORE2026 A*", "USENIX_Security_CORE_Astar.png"),
        ("ACL — ICORE2026 A*", "ACL_CORE_Astar.png"),
        ("ICSE — ICORE2026 A*", "ICSE_CORE_Astar.png"),
        ("FSE — ICORE2026 A*", "FSE_CORE_Astar.png"),
        ("ICPC — ICORE2026 A", "ICPC_CORE_A.png"),
    ]
    for idx, (caption, fn) in enumerate(evidence):
        if idx:
            doc.add_page_break()
        add_heading(doc, caption, 2)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_picture_with_alt(
            p.add_run(), ROOT / "evidence" / "core" / fn, Inches(6.75),
            f"Official ICORE Conference Portal screenshot: {caption}."
        )
        c = doc.add_paragraph("Official ICORE Conference Portal search captured 17 October 2026. Source: https://portal.core.edu.au/conf-ranks/")
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.runs[0].italic = True
        c.runs[0].font.size = Pt(8)

    # Document properties
    props = doc.core_properties
    props.title = "GateKeeper AI Phase 1"
    props.subject = "G3: Empowering CI/CD with Agentic AI for Enhanced Software Security"
    props.author = "Anas Imran, Hamza Farooq, Maaz Bhatti"
    props.keywords = "agentic AI, CI/CD, SAST, vulnerability triage, software security"
    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    build()
